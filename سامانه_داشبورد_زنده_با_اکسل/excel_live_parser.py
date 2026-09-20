import os
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

def parse_excel_claims(excel_path: str) -> dict:
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # Locate target sheet
    sheet_name = 'TOTAL'
    if sheet_name not in wb.sheetnames:
        for s in wb.sheetnames:
            if 'TOTAL' in s.upper():
                sheet_name = s
                break
        else:
            sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    station_names_map = {
        "Z10-5": "ایستگاه دریاچه",
        "Z10-7 1": "ایستگاه پژوهش",
        "Z10-2": "ایستگاه اتریش",
        "Z10": "ایستگاه دهکده المپیک",
        "X10": "ایستگاه جنت‌آباد",
        "W10": "ایستگاه ایرانشهر",
        "V10": "ایستگاه سردار جنگل",
        "Z10-9": "کارگاه پایانه (وردآورد)",
        "SA": "سگمنت علی‌آباد",
        "TURL10": "کلیات خط ۱۰ و تونل",
        "TBM1": "دستگاه حفار مکانیزه TBM1",
        "Contract 2nd": "امور قراردادهای پیمانکاران دست‌دوم"
    }

    employer_claims = []
    subcontractor_claims = []
    current_station_code = None
    current_station_name = None
    is_in_subcontractors = False
    is_in_others = False

    def to_float(val):
        try:
            if val is None or val == "" or val == "-":
                return 0.0
            return float(val)
        except Exception:
            return 0.0

    # Start from row 3 (row 1 is title header, row 2 is sheet summary totals)
    for r in range(3, ws.max_row + 1):
        col2 = ws.cell(r, 2).value
        col3 = ws.cell(r, 3).value
        col4 = ws.cell(r, 4).value
        col5 = ws.cell(r, 5).value
        col6 = ws.cell(r, 6).value
        col7 = ws.cell(r, 7).value
        col8 = ws.cell(r, 8).value
        col9 = ws.cell(r, 9).value
        col10 = ws.cell(r, 10).value
        col11 = ws.cell(r, 11).value

        col2_str = str(col2).strip() if col2 is not None else ""
        col3_str = str(col3).strip() if col3 is not None else ""

        # Skip completely empty rows
        if not col2_str and not col3_str and not col8 and not col9 and not col10:
            continue

        # Check section boundaries
        if col2_str == "OTHERS" or "سایر اقدامات" in col3_str:
            is_in_others = True
            continue

        if is_in_others:
            continue

        if col2_str == "Contract 2nd" or "پیمانکاران دست دوم" in col3_str or "امورقراردادهای پیمانکاران" in col3_str:
            is_in_subcontractors = True
            current_station_code = "Contract 2nd"
            current_station_name = "امور قراردادهای پیمانکاران دست‌دوم"
            continue

        # Check if row is a Station Header row
        if (col2_str in station_names_map or (col2_str and not '.' in col2_str and not col8 and not col9)) and not is_in_subcontractors:
            current_station_code = col2_str
            current_station_name = station_names_map.get(col2_str, col3_str or col2_str)
            continue

        # Parse claim row
        code_str = col2_str if col2_str else "-"
        title_str = col3_str if col3_str else "بدون عنوان"
        circ_str = str(col4).strip() if col4 is not None else ""
        actions_str = str(col5).strip() if col5 is not None else ""
        method_str = str(col6).strip() if col6 is not None else ""
        in_prog_str = str(col7).strip() if col7 is not None else ""

        if is_in_subcontractors:
            claimed_val = to_float(col8)
            saved_val = to_float(col9)
            approved_val = max(0.0, claimed_val - saved_val)
            rate = round((saved_val / claimed_val) * 100, 1) if claimed_val > 0 else 0.0

            sub_item = {
                "row": r,
                "code": code_str,
                "contractor": title_str,
                "title": title_str,
                "claimed_amount": claimed_val,
                "approved_amount": approved_val,
                "saved_amount": saved_val,
                "success_rate": rate,
                "circulars": circ_str,
                "actions": actions_str,
                "realization_method": method_str,
                "realization_method_excel": method_str,
                "has_financial": (claimed_val > 0 or saved_val > 0 or approved_val > 0)
            }
            subcontractor_claims.append(sub_item)
        else:
            init_val = to_float(col8)
            rec_val = to_float(col9)
            pot_val = to_float(col10)
            st_name = current_station_name if current_station_name else "کلیات پروژه"

            emp_item = {
                "row": r,
                "code": code_str,
                "station_code": current_station_code or "-",
                "station_name": st_name,
                "title": title_str,
                "circulars": circ_str,
                "actions": actions_str,
                "in_progress": in_prog_str,
                "initial_claim": init_val,
                "received_amount": rec_val,
                "potential_amount": pot_val,
                "realization_method": method_str,
                "realization_method_excel": method_str,
                "has_financial": (init_val > 0 or rec_val > 0 or pot_val > 0)
            }
            employer_claims.append(emp_item)

    # Compute Summaries
    tot_emp_initial = sum(c["initial_claim"] for c in employer_claims)
    tot_emp_received = sum(c["received_amount"] for c in employer_claims)
    tot_emp_potential = sum(c["potential_amount"] for c in employer_claims)
    tot_emp_pending = max(0.0, tot_emp_initial - tot_emp_received)
    overall_emp_rate = round((tot_emp_received / tot_emp_initial) * 100, 2) if tot_emp_initial > 0 else 0.0

    tot_sub_claimed = sum(s["claimed_amount"] for s in subcontractor_claims)
    tot_sub_saved = sum(s["saved_amount"] for s in subcontractor_claims)
    tot_sub_approved = sum(s["approved_amount"] for s in subcontractor_claims)
    overall_sub_rate = round((tot_sub_saved / tot_sub_claimed) * 100, 1) if tot_sub_claimed > 0 else 0.0

    tot_realized_all = tot_emp_received + tot_sub_saved
    tot_potential_all = tot_emp_potential + tot_sub_saved

    summary = {
        "employer_totals": {
            "initial_claimed_rial": tot_emp_initial,
            "received_to_date_rial": tot_emp_received,
            "potential_all_stations_rial": tot_emp_potential,
            "claims_count": len(employer_claims)
        },
        "subcontractor_defense_totals": {
            "claimed_by_subs_rial": tot_sub_claimed,
            "approved_to_pay_rial": tot_sub_approved,
            "saved_defended_rial": tot_sub_saved,
            "defense_success_rate_pct": overall_sub_rate,
            "claims_count": len(subcontractor_claims)
        },
        "grand_value_creation": {
            "realized_value_to_date_rial": tot_realized_all,
            "grand_potential_rial": tot_potential_all,
            "total_active_cases": len(employer_claims) + len(subcontractor_claims)
        },
        "employer": {
            "total_initial_claim": tot_emp_initial,
            "total_received": tot_emp_received,
            "total_pending": tot_emp_pending,
            "total_potential": tot_emp_potential,
            "realization_rate": overall_emp_rate,
            "claims_count": len(employer_claims)
        },
        "subcontractor": {
            "total_claimed": tot_sub_claimed,
            "total_saved": tot_sub_saved,
            "total_approved": tot_sub_approved,
            "success_rate": overall_sub_rate,
            "claims_count": len(subcontractor_claims)
        },
        "combined": {
            "total_financial_impact": tot_realized_all,
            "total_potential_value": tot_potential_all,
            "total_active_cases": len(employer_claims) + len(subcontractor_claims)
        }
    }

    return {
        "summary": summary,
        "employer_claims": employer_claims,
        "subcontractor_claims": subcontractor_claims
    }

if __name__ == "__main__":
    p = r"c:\Users\Hooman\Desktop\داشبورد امور قرارداد\میزان کلیم\claim metro L100.xlsx"
    res = parse_excel_claims(p)
    print("Parsed Successfully!")
    print(f"Employer claims: {len(res['employer_claims'])}")
    print(f"Subcontractor claims: {len(res['subcontractor_claims'])}")
    print("Summary:")
    print(json.dumps(res['summary']['grand_value_creation'], ensure_ascii=False, indent=2))
